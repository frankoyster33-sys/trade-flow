"""Read-only extraction. Never saves or modifies the input workbook."""
import sys,json,csv,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent.parent/'.python'))
file=Path(sys.argv[1]); ext=file.suffix.lower()
rows=[]
if ext=='.xlsx':
 import openpyxl
 with zipfile.ZipFile(file) as z:
  if sum(x.file_size for x in z.infolist())>100*1024*1024: raise ValueError('Excel 解压后超过100MB，请拆分文件')
 wb=openpyxl.load_workbook(file,read_only=True,data_only=True)
 for ws in wb.worksheets:
  if (ws.max_row or 0)>2500 or (ws.max_column or 0)>100: raise ValueError('单表超过2500行或100列，请拆分询价')
  for i,row in enumerate(ws.iter_rows()):
   if i>=2500 or len(row)>100: raise ValueError('单表超过2500行或100列，请拆分询价')
   cells=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
   if cells: rows.append(f'{ws.title}: '+' | '.join(cells))
 wb.close()
elif ext=='.xls':
 import pandas as pd
 book=pd.read_excel(file,sheet_name=None,header=None,engine='xlrd')
 for name,frame in book.items():
  if frame.shape[0]>2500 or frame.shape[1]>100: raise ValueError('表格过大，请拆分询价')
  for i,row in frame.fillna('').iterrows(): rows.append(f'{name} 行{i+1}: '+' | '.join(map(str,row)))
elif ext=='.csv':
 text=file.read_bytes()
 try: content=text.decode('utf-8-sig')
 except UnicodeDecodeError: content=text.decode('gb18030')
 for i,row in enumerate(csv.reader(content.splitlines())):
  if i>2500: raise ValueError('表格超过2500行，请拆分询价')
  rows.append(f'行{i+1}: '+' | '.join(row))
else: raise ValueError('不支持的表格类型')
out='\n'.join(rows)
if len(out)>100000: raise ValueError('询价表格内容过多，请拆分后导入')
print(out)
